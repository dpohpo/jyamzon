#!/usr/bin/env python3
"""Verify jyamzon JSON and Excel outputs are internally consistent."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BROAD_DATA_DIR = Path("data/amazon_3c/bsr_new_releases")
FILTERED_DATA_DIR = Path("data/amazon_3c/filtered_30_opportunities")
BROAD_XLSX = Path("outputs/amazon_3c_bsr_new_releases/amazon_3c_new_releases_bsr_crawl.xlsx")
FILTERED_XLSX = Path("outputs/amazon_3c_filtered_opportunities/amazon_3c_bsr_filtered_30_opportunities.xlsx")


def load_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def count_excel_asins(path: Path, sheet_name: str) -> dict[str, Any]:
    if not path.exists():
        return {"exists": False, "rows": 0, "unique_asins": 0, "images": 0, "headers": []}
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    asin_col = headers.index("ASIN") + 1 if "ASIN" in headers else None
    asins = []
    if asin_col:
        asins = [ws.cell(row, asin_col).value for row in range(2, ws.max_row + 1) if ws.cell(row, asin_col).value]
    return {
        "exists": True,
        "rows": len(asins),
        "unique_asins": len(set(asins)),
        "duplicate_asins": sorted(asin for asin, count in Counter(asins).items() if count > 1),
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "images": len(getattr(ws, "_images", [])),
        "headers": headers,
        "sheets": wb.sheetnames,
    }


def verify_broad() -> dict[str, Any]:
    summary = load_json(BROAD_DATA_DIR / "run_summary.json", {})
    excel = count_excel_asins(BROAD_XLSX, "3C新品榜单")
    category_picks = int(summary.get("category_picks", 0) or 0)
    products_with_title = int(summary.get("products_with_title", 0) or 0)
    coverage = round(products_with_title / category_picks * 100, 1) if category_picks else 0.0
    return {
        "summary": summary,
        "excel": excel,
        "coverage_percent": coverage,
        "ok": bool(excel.get("exists")) and excel.get("unique_asins") == category_picks,
    }


def verify_filtered(target: int) -> dict[str, Any]:
    summary = load_json(FILTERED_DATA_DIR / "summary.json", {})
    selected = load_json(FILTERED_DATA_DIR / "selected_30.json", [])
    curated = load_json(FILTERED_DATA_DIR / "selected_30_curated.json", [])
    excel = count_excel_asins(FILTERED_XLSX, "30个筛选机会品")
    selected_count = len({row.get("asin") for row in selected if row.get("asin")})
    curated_count = len({row.get("asin") for row in curated if row.get("asin")})
    excel_count = int(excel.get("unique_asins") or 0)
    counts_match = selected_count == curated_count == excel_count
    return {
        "summary": summary,
        "selected_json_unique_asins": selected_count,
        "curated_json_unique_asins": curated_count,
        "excel": excel,
        "counts_match": counts_match,
        "target_met": excel_count >= target,
        "ok": counts_match and excel_count >= target and not excel.get("duplicate_asins"),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--target", type=int, default=30, help="Expected filtered workbook ASIN count.")
    parser.add_argument("--filtered-only", action="store_true", help="Verify only filtered shortlist outputs.")
    args = parser.parse_args()

    result = {"filtered": verify_filtered(args.target)}
    if not args.filtered_only:
        result["broad"] = verify_broad()
    result["ok"] = bool(result["filtered"]["ok"] and result.get("broad", {"ok": True})["ok"])
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["ok"] else 2


if __name__ == "__main__":
    raise SystemExit(main())

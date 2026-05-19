#!/usr/bin/env python3
"""Build a polished Excel workbook from the Amazon 3C BSR crawl output."""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image as PILImage


DATA_DIR = Path("data/amazon_3c/bsr_new_releases")
OUTPUT_DIR = Path("outputs/amazon_3c_bsr_new_releases")
OUTPUT_XLSX = OUTPUT_DIR / "amazon_3c_new_releases_bsr_crawl.xlsx"
THUMB_DIR = DATA_DIR / "thumbs"


MAIN_HEADERS = [
    "序号",
    "图片",
    "榜单",
    "一级大类",
    "类目路径",
    "类目名",
    "类目节点",
    "榜内排名",
    "ASIN",
    "标题",
    "品牌",
    "价格",
    "价格数值",
    "评分",
    "评论数",
    "相关性标记",
    "抓取备注",
    "Bullet 1",
    "Bullet 2",
    "Bullet 3",
    "Bullet 4",
    "Bullet 5",
    "图片URL",
    "本地图片",
    "商品链接",
    "榜单链接",
    "BSR文本",
]


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def parse_price(value: str) -> float | None:
    if not value:
        return None
    match = re.search(r"[\d,]+(?:\.\d+)?", value)
    if not match:
        return None
    return float(match.group(0).replace(",", ""))


def parse_int(value: str) -> int | None:
    if value is None or value == "":
        return None
    match = re.search(r"[\d,]+", str(value))
    if not match:
        return None
    return int(match.group(0).replace(",", ""))


def parse_float(value: str) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def classify(row: dict[str, Any]) -> str:
    path = (row.get("category_path") or "").lower()
    title = (row.get("title") or "").lower()
    service_terms = ["service", "warrant", "license", "installation", "coverage"]
    software_terms = ["mac games", "pc games", "edition", "nintendo switch", "playstation", "xbox"]
    light_terms = [
        "accessories",
        "cases",
        "cables",
        "cord management",
        "mounts",
        "headphones",
        "microphones",
        "power strips",
        "surge protectors",
        "tablet accessories",
        "laptop accessories",
        "bags & cases",
    ]
    heavy_terms = ["laptops", "desktops", "servers", "televisions", "projectors", "cameras", "tablets"]
    if any(term in path or term in title for term in service_terms):
        return "排除/谨慎：服务、授权或保修"
    if "video games" in path and any(term in path or term in title for term in software_terms):
        return "低相关：游戏软件/平台内容"
    if any(term in path for term in light_terms):
        return "优先看：轻小件/配件"
    if any(term in path for term in heavy_terms):
        return "谨慎：整机/高客单/售后重"
    return "可观察：3C硬件/周边"


def crawl_note(row: dict[str, Any]) -> str:
    notes: list[str] = []
    if not row.get("title"):
        notes.append("标题缺失")
    if not row.get("price"):
        notes.append("价格缺失")
    if not row.get("review_count"):
        notes.append("评论数为空或新品无评论")
    if not row.get("bullet_5"):
        notes.append("五点不足5条")
    if not row.get("local_image"):
        notes.append("图片未下载")
    if not notes:
        return "字段完整"
    return "；".join(notes)


def make_thumb(image_path: str, asin: str) -> str:
    if not image_path:
        return ""
    src = Path(image_path)
    if not src.exists() or src.stat().st_size == 0:
        return ""
    THUMB_DIR.mkdir(parents=True, exist_ok=True)
    out = THUMB_DIR / f"{asin}.png"
    try:
        with PILImage.open(src) as image:
            image = image.convert("RGB")
            image.thumbnail((96, 96))
            canvas = PILImage.new("RGB", (96, 96), "white")
            x = (96 - image.width) // 2
            y = (96 - image.height) // 2
            canvas.paste(image, (x, y))
            canvas.save(out, "PNG", optimize=True)
        return str(out)
    except Exception:
        return ""


def style_header(ws, row: int, fill: str = "1F4E78") -> None:
    for cell in ws[row]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_border(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    side = Side(style="thin", color="D9E2F3")
    border = Border(left=side, right=side, top=side, bottom=side)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_main_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.active
    ws.title = "3C新品榜单"
    ws.append(MAIN_HEADERS)
    style_header(ws, 1)
    ws.freeze_panes = "J2"

    for index, row in enumerate(rows, start=1):
        price_num = parse_price(row.get("price", ""))
        rating_num = parse_float(row.get("rating", ""))
        review_num = parse_int(row.get("review_count", ""))
        relevance = classify(row)
        note = crawl_note(row)
        data = [
            index,
            "",
            row.get("source", ""),
            row.get("department", ""),
            row.get("category_path", ""),
            row.get("category_name", ""),
            row.get("category_node", ""),
            row.get("rank_on_page", ""),
            row.get("asin", ""),
            row.get("title", ""),
            row.get("brand", ""),
            row.get("price", ""),
            price_num,
            rating_num,
            review_num,
            relevance,
            note,
            row.get("bullet_1", ""),
            row.get("bullet_2", ""),
            row.get("bullet_3", ""),
            row.get("bullet_4", ""),
            row.get("bullet_5", ""),
            row.get("image_url", ""),
            row.get("local_image", ""),
            row.get("product_url", ""),
            row.get("category_url", ""),
            row.get("bsr_text", ""),
        ]
        ws.append(data)
        excel_row = index + 1
        ws.row_dimensions[excel_row].height = 76
        thumb = make_thumb(row.get("local_image", ""), row.get("asin", f"row{index}"))
        if thumb:
            image = XLImage(thumb)
            image.width = 72
            image.height = 72
            ws.add_image(image, f"B{excel_row}")
        for col in [23, 24, 25, 26]:
            cell = ws.cell(excel_row, col)
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

        flag_cell = ws.cell(excel_row, 16)
        if str(flag_cell.value).startswith("优先"):
            flag_cell.fill = PatternFill("solid", fgColor="E2F0D9")
        elif str(flag_cell.value).startswith("排除"):
            flag_cell.fill = PatternFill("solid", fgColor="F4CCCC")
        elif str(flag_cell.value).startswith("低相关"):
            flag_cell.fill = PatternFill("solid", fgColor="FFF2CC")
        elif str(flag_cell.value).startswith("谨慎"):
            flag_cell.fill = PatternFill("solid", fgColor="FCE4D6")
        else:
            flag_cell.fill = PatternFill("solid", fgColor="DDEBF7")

    widths = {
        "A": 7,
        "B": 14,
        "C": 14,
        "D": 22,
        "E": 42,
        "F": 24,
        "G": 14,
        "H": 10,
        "I": 14,
        "J": 52,
        "K": 24,
        "L": 12,
        "M": 11,
        "N": 9,
        "O": 10,
        "P": 24,
        "Q": 24,
        "R": 45,
        "S": 45,
        "T": 45,
        "U": 45,
        "V": 45,
        "W": 30,
        "X": 28,
        "Y": 30,
        "Z": 30,
        "AA": 44,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    set_border(ws, 1, len(rows) + 1, 1, len(MAIN_HEADERS))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(MAIN_HEADERS))}{len(rows) + 1}"

    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=1, max_col=len(MAIN_HEADERS)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in ["A", "H", "I", "L", "M", "N", "O"]:
        for cell in ws[col]:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    table_ref = f"A1:{get_column_letter(len(MAIN_HEADERS))}{len(rows) + 1}"
    table = Table(displayName="Amazon3CNewReleases", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def build_dashboard(wb: Workbook, rows: list[dict[str, Any]], summary: dict[str, Any]) -> None:
    ws = wb.create_sheet("概览", 0)
    ws["A1"] = "Amazon US 3C New Releases / BSR 爬虫结果"
    ws["A1"].font = Font(size=18, bold=True, color="1F4E78")
    ws.merge_cells("A1:H1")
    ws["A2"] = f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}；数据源：Amazon New Releases 榜单页 + 商品详情页；未调用 MCP。"
    ws.merge_cells("A2:H2")

    metrics = [
        ("类目数", summary.get("category_picks", len(rows))),
        ("成功抓到标题", summary.get("products_with_title", "")),
        ("成功抓到价格", summary.get("products_with_price", "")),
        ("成功抓到图片", summary.get("products_with_image", "")),
        ("有评论数字段", summary.get("products_with_review_count", "")),
        ("有至少5条Bullet", summary.get("products_with_five_bullets", "")),
    ]
    for idx, (label, value) in enumerate(metrics):
        col = 1 + (idx % 3) * 3
        row = 4 + (idx // 3) * 3
        ws.cell(row, col, label)
        ws.cell(row + 1, col, value)
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        ws.cell(row, col).fill = PatternFill("solid", fgColor="D9EAF7")
        ws.cell(row, col).font = Font(bold=True, color="1F4E78")
        ws.cell(row + 1, col).font = Font(size=16, bold=True)
        ws.cell(row + 1, col).alignment = Alignment(horizontal="center")

    ws["A11"] = "相关性/风险标记统计"
    ws["A11"].font = Font(size=13, bold=True, color="1F4E78")
    flags = Counter(classify(row) for row in rows)
    ws.append([])
    ws.append(["标记", "数量", "含义"])
    style_header(ws, 13, "5B9BD5")
    meanings = {
        "优先看：轻小件/配件": "更接近轻资产3C配件选品，可优先二次筛选",
        "可观察：3C硬件/周边": "属于3C硬件/周边，但还需看利润、认证和售后",
        "谨慎：整机/高客单/售后重": "整机或高客单类，供应链、退货、售后压力更大",
        "低相关：游戏软件/平台内容": "偏内容/软件，不是典型实体3C私标品",
        "排除/谨慎：服务、授权或保修": "服务/授权/保修，不适合作为实体产品选品",
    }
    r = 14
    for flag, count in flags.most_common():
        ws.cell(r, 1, flag)
        ws.cell(r, 2, count)
        ws.cell(r, 3, meanings.get(flag, ""))
        r += 1

    ws["E11"] = "一级大类覆盖"
    ws["E11"].font = Font(size=13, bold=True, color="1F4E78")
    ws["E13"] = "一级大类"
    ws["F13"] = "数量"
    style_header(ws, 13, "70AD47")
    dept_counts = Counter(row.get("department", "") for row in rows)
    r = 14
    for dept, count in dept_counts.most_common():
        ws.cell(r, 5, dept)
        ws.cell(r, 6, count)
        r += 1

    ws["A23"] = "使用建议"
    ws["A23"].font = Font(size=13, bold=True, color="1F4E78")
    notes = [
        "先筛选“优先看：轻小件/配件”，再排除评论数过高、价格过低或服务/授权类。",
        "New Releases 榜单只说明新品热度，不等于销量或利润；后续仍需验证搜索量、购买率、PPC、集中度。",
        "评论数为空可能是新品无评论，也可能是详情页字段未公开；不要直接当作低竞争。",
        "图片已嵌入主表；原图 URL 和本地图片路径也保留，便于复核。",
    ]
    for idx, note in enumerate(notes, start=24):
        ws.cell(idx, 1, f"{idx - 23}. {note}")
        ws.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=8)

    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["E"].width = 28
    set_border(ws, 13, max(14 + len(flags), 14 + len(dept_counts), 27), 1, 8)


def build_logs_sheet(wb: Workbook, logs: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("抓取日志")
    headers = ["URL", "榜单", "一级大类", "类目路径", "深度", "状态", "页面ASIN数", "子类目数"]
    ws.append(headers)
    style_header(ws, 1, "7F7F7F")
    for log in logs:
        ws.append(
            [
                log.get("url", ""),
                log.get("source", ""),
                log.get("department", ""),
                log.get("path", ""),
                log.get("depth", ""),
                log.get("status", ""),
                log.get("asins", ""),
                log.get("children", ""),
            ]
        )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(logs) + 1}"
    widths = [60, 14, 24, 48, 8, 12, 12, 12]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    set_border(ws, 1, len(logs) + 1, 1, len(headers))


def build_field_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("字段覆盖")
    fields = [
        ("标题", "title"),
        ("价格", "price"),
        ("评分", "rating"),
        ("评论数", "review_count"),
        ("Bullet 1", "bullet_1"),
        ("Bullet 5", "bullet_5"),
        ("图片URL", "image_url"),
        ("本地图片", "local_image"),
        ("BSR文本", "bsr_text"),
    ]
    ws.append(["字段", "非空数量", "覆盖率", "说明"])
    style_header(ws, 1, "8064A2")
    total = len(rows)
    explanations = {
        "BSR文本": "Amazon 详情页不一定暴露或页面结构变化，本次保留空值",
        "评论数": "新品可能无评论；空值不等于0",
        "Bullet 5": "有些页面少于5条五点或页面结构不同",
    }
    for label, key in fields:
        count = sum(1 for row in rows if row.get(key))
        ws.append([label, count, count / total if total else 0, explanations.get(label, "")])
    for cell in ws["C"][1:]:
        cell.number_format = "0.0%"
    for col, width in {"A": 18, "B": 12, "C": 12, "D": 60}.items():
        ws.column_dimensions[col].width = width
    set_border(ws, 1, len(fields) + 1, 1, 4)


def save_and_verify(path: Path) -> dict[str, Any]:
    wb = load_workbook(path)
    checks = {
        "sheets": wb.sheetnames,
        "main_rows": wb["3C新品榜单"].max_row - 1,
        "main_cols": wb["3C新品榜单"].max_column,
        "dashboard_rows": wb["概览"].max_row,
    }
    required = {"概览", "3C新品榜单", "字段覆盖", "抓取日志"}
    checks["required_sheets_ok"] = required.issubset(set(wb.sheetnames))
    checks["nonempty_main_ok"] = checks["main_rows"] >= 50
    return checks


def main() -> int:
    rows = load_json(DATA_DIR / "products.json")
    summary = load_json(DATA_DIR / "run_summary.json")
    logs = load_json(DATA_DIR / "category_page_logs.json")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    build_main_sheet(wb, rows)
    build_dashboard(wb, rows, summary)
    build_field_sheet(wb, rows)
    build_logs_sheet(wb, logs)

    wb.save(OUTPUT_XLSX)
    checks = save_and_verify(OUTPUT_XLSX)
    (OUTPUT_DIR / "verification.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT_XLSX), **checks}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""Collect granular SellerSprite ASIN and keyword data through OpenCLI.

This script uses an already logged-in Chrome profile controlled by OpenCLI. It
keeps the Excel output normalized: each cell contains one metric, one flag, or
one identifier. Raw visible DOM text is preserved only in the JSON audit file.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


OPENCLI_PACKAGE = "@jackwener/opencli@1.7.22"
AMAZON_SESSION = "amazonss"
SELLERSPRITE_SESSION = "sellersprite"
SELLERSPRITE_REVERSING_URL = "https://www.sellersprite.com/v3/reversing"
SELLERSPRITE_KEYWORD_REVERSE_URL = "https://www.sellersprite.com/v3/keyword-reverse"
SELLERSPRITE_KEYWORD_MINER_URL = "https://www.sellersprite.com/v3/keyword-miner"

PLUGIN_SELECTORS = {
    "matrix_meta": "#__ss_matrix_container_meta",
    "quick_view_page": "#seller-sprite-extension-quick-view-listing-page",
    "quick_view_detail": "#seller-sprite-extension-quick-view-listing",
    "inventory": "#sellersprite-extension-inventory",
    "inventory_surplus": "#sellersprite-extension-Inventory-surplus-count",
    "main_relation": "#seller-sprite-extension-main-relation",
    "listing_buttons": "#seller-sprite-listing-btn-group-box",
    "product_store_button": "#seller-sprite-extension-add-product-store-btn",
    "find_similar": "#seller-sprite-extension-find-similar-listing",
}


def run_opencli(session: str, *args: str, timeout: int = 60) -> str:
    cmd = ["npx", "-y", OPENCLI_PACKAGE, "browser", session, *args]
    result = subprocess.run(
        cmd,
        check=False,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        timeout=timeout,
    )
    if result.returncode != 0:
        raise RuntimeError(
            f"OpenCLI failed: {' '.join(cmd)}\nSTDOUT:\n{result.stdout}\nSTDERR:\n{result.stderr}"
        )
    return result.stdout.strip()


def eval_js(session: str, js: str, timeout: int = 60) -> Any:
    out = run_opencli(session, "eval", js, timeout=timeout)
    try:
        return json.loads(out)
    except json.JSONDecodeError:
        return out


def clean_text(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def parse_int(value: str | int | float | None) -> int | None:
    if value is None:
        return None
    text = str(value)
    match = re.search(r"-?[\d,]+", text)
    if not match:
        return None
    return int(match.group(0).replace(",", ""))


def parse_float(value: str | int | float | None) -> float | None:
    if value is None:
        return None
    text = str(value).replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    return float(match.group(0))


def parse_percent(value: str | None) -> float | None:
    if not value or "%" not in value:
        return None
    number = parse_float(value)
    return number / 100 if number is not None else None


def parse_money(value: str | None) -> tuple[str, float | None]:
    if not value:
        return "", None
    currency = "USD" if "$" in value else ""
    return currency, parse_float(value)


def safe_slug(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_")
    return slug or "item"


def infer_asin(value: str) -> str:
    match = re.search(r"(?:/dp/|/gp/product/|^)([A-Z0-9]{10})(?:[/?#]|$)", value)
    if match:
        return match.group(1)
    match = re.search(r"\b([A-Z0-9]{10})\b", value)
    if match:
        return match.group(1)
    raise ValueError(f"Could not infer ASIN from: {value}")


@dataclass
class MetricRow:
    asin: str
    field: str
    value: Any
    unit: str
    source: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "asin": self.asin,
            "field": self.field,
            "value": self.value,
            "unit": self.unit,
            "source": self.source,
        }


def add_metric(rows: list[MetricRow], asin: str, field: str, value: Any, unit: str, source: str) -> None:
    if value is None or value == "":
        return
    rows.append(MetricRow(asin=asin, field=field, value=value, unit=unit, source=source))


def add_regex_metric(
    rows: list[MetricRow],
    asin: str,
    text: str,
    field: str,
    pattern: str,
    source: str,
    parser=lambda x: x,
    unit: str = "",
) -> None:
    match = re.search(pattern, text)
    if match:
        add_metric(rows, asin, field, parser(match.group(1).strip()), unit, source)


def parse_launch_date(value: str) -> tuple[str, int | None]:
    date_match = re.search(r"\d{4}-\d{2}-\d{2}", value)
    days_match = re.search(r"\((\d+)\s*天\)", value)
    return (date_match.group(0) if date_match else "", parse_int(days_match.group(1)) if days_match else None)


def parse_summary_metrics(asin: str, blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_name = {block.get("name"): clean_text(block.get("text", "")) for block in blocks}
    quick = by_name.get("quick_view_page", "")
    detail = by_name.get("quick_view_detail", "")
    rows: list[MetricRow] = []

    add_regex_metric(rows, asin, quick, "quality_score", r"质量得分\s*([0-9.]+)", "quick_view_page", parse_float)
    parent_sales = re.search(r"近30天销量\(父体\)\s*([0-9,]+)(?:\s*\(([0-9,]+)\))?", quick)
    if parent_sales:
        add_metric(rows, asin, "parent_sales_30d", parse_int(parent_sales.group(1)), "units", "quick_view_page")
        add_metric(
            rows,
            asin,
            "parent_sales_30d_parenthetical",
            parse_int(parent_sales.group(2)),
            "units",
            "quick_view_page",
        )
    for field, pattern in [
        ("listing_revenue_30d", r"Listing销售额\s*(\$[0-9,]+(?:\.[0-9]+)?)"),
        ("average_price", r"均价\s*(\$[0-9,]+(?:\.[0-9]+)?)"),
        ("fba_fee", r"FBA\s*费用\s*(\$[0-9,]+(?:\.[0-9]+)?)"),
    ]:
        match = re.search(pattern, quick)
        if match:
            currency, amount = parse_money(match.group(1))
            add_metric(rows, asin, field, amount, currency, "quick_view_page")
    add_regex_metric(rows, asin, quick, "bsr_main_rank", r"BSR\s*([0-9,]+)", "quick_view_page", parse_int)
    add_regex_metric(rows, asin, quick, "variation_count", r"变体数\s*([0-9,]+)", "quick_view_page", parse_int)
    launch = re.search(r"上架时间\s*([0-9-]+(?:\(\d+\s*天\))?)", quick)
    if launch:
        launch_date, days_live = parse_launch_date(launch.group(1))
        add_metric(rows, asin, "launch_date", launch_date, "date", "quick_view_page")
        add_metric(rows, asin, "days_live", days_live, "days", "quick_view_page")

    add_regex_metric(rows, asin, detail, "asin", r"ASIN[:：]\s*([A-Z0-9]{10})", "quick_view_detail")
    add_regex_metric(rows, asin, detail, "brand", r"品牌[:：]\s*(.*?)\s+卖家[:：]", "quick_view_detail")
    add_regex_metric(rows, asin, detail, "seller", r"卖家[:：]\s*(.*?)\s+配送[:：]", "quick_view_detail")
    add_regex_metric(rows, asin, detail, "fulfillment", r"配送[:：]\s*([A-Z]+)", "quick_view_detail")
    add_regex_metric(rows, asin, detail, "seller_count", r"配送[:：]\s*[A-Z]+\s*卖家[:：]\s*([0-9,]+)", "quick_view_detail", parse_int)
    add_regex_metric(
        rows,
        asin,
        detail,
        "parent_sales_30d_detail",
        r"近30天销量\(父体\)[:：]\s*([0-9,]+)",
        "quick_view_detail",
        parse_int,
        "units",
    )
    add_regex_metric(
        rows,
        asin,
        detail,
        "child_sales_30d",
        r"近30天销量\(子体\)[:：]\s*([0-9,+]+)",
        "quick_view_detail",
        lambda x: x,
        "units",
    )
    for field, pattern in [
        ("sales_amount_30d", r"销售额[:：]\s*(\$[0-9,]+(?:\.[0-9]+)?)"),
        ("fba_fee_detail", r"FBA费用[:：]\s*FBA\s*费用[:：]?\s*(\$[0-9,]+(?:\.[0-9]+)?)"),
        ("price", r"价格[:：]\s*(\$[0-9,]+(?:\.[0-9]+)?)"),
    ]:
        match = re.search(pattern, detail)
        if match:
            currency, amount = parse_money(match.group(1))
            add_metric(rows, asin, field, amount, currency, "quick_view_detail")
    add_regex_metric(
        rows,
        asin,
        detail,
        "gross_margin",
        r"毛利率[:：]\s*([0-9.]+%)",
        "quick_view_detail",
        parse_percent,
        "ratio",
    )
    add_regex_metric(rows, asin, detail, "variation_count_detail", r"变体数[:：]\s*([0-9,]+)", "quick_view_detail", parse_int)
    rating = re.search(r"评分\(评分数\)[:：]\s*([0-9.]+)\(([0-9,]+)\)", detail)
    if rating:
        add_metric(rows, asin, "rating", parse_float(rating.group(1)), "stars", "quick_view_detail")
        add_metric(rows, asin, "rating_count", parse_int(rating.group(2)), "reviews", "quick_view_detail")
    add_regex_metric(rows, asin, detail, "color", r"Color[:：]\s*(.*?)\s+颜色[:：]", "quick_view_detail")
    add_regex_metric(rows, asin, detail, "size", r"Size[:：]\s*(.*?)\s+尺寸[:：]", "quick_view_detail")
    add_regex_metric(rows, asin, detail, "product_weight", r"商品重量[:：]\s*(.*?)\s+商品尺寸[:：]", "quick_view_detail")
    add_regex_metric(rows, asin, detail, "product_dimensions", r"商品尺寸[:：]\s*(.*?)\s+包装重量[:：]", "quick_view_detail")
    add_regex_metric(rows, asin, detail, "package_weight", r"包装重量[:：]\s*(.*?)\s+包装尺寸[:：]", "quick_view_detail")
    add_regex_metric(rows, asin, detail, "package_dimensions", r"包装尺寸[:：]\s*(.*?)\s+上架时间[:：]", "quick_view_detail")
    detail_launch = re.search(r"上架时间[:：]\s*([0-9-]+)\s*\((\d+)天\)", detail)
    if detail_launch:
        add_metric(rows, asin, "launch_date_detail", detail_launch.group(1), "date", "quick_view_detail")
        add_metric(rows, asin, "days_live_detail", parse_int(detail_launch.group(2)), "days", "quick_view_detail")
    for field, pattern in [
        ("all_traffic_keywords", r"全部流量词[:：]\s*([0-9,]+)"),
        ("natural_search_keywords", r"自然搜索词[:：]\s*([0-9,]+)"),
        ("ad_traffic_keywords", r"广告流量词[:：]\s*([0-9,]+)"),
        ("search_recommendation_keywords", r"搜索推荐词[:：]\s*([0-9,]+)"),
    ]:
        add_regex_metric(rows, asin, detail, field, pattern, "quick_view_detail", parse_int, "keywords")

    return [row.as_dict() for row in rows]


def parse_rank_entry(text: str) -> dict[str, Any]:
    text = clean_text(text)
    if not text or "前3页无排名" in text:
        return {"rank": None, "page": None, "position": None, "page_total": None, "observed_at": "", "not_top_3_pages": True}
    match = re.search(
        r"(?P<rank>\d+)\s+第(?P<page>\d+)页,\s*(?P<position>\d+)\s*/\s*(?P<total>\d+)\s+(?P<observed_at>[^ ]+排名)",
        text,
    )
    if not match:
        return {"rank": None, "page": None, "position": None, "page_total": None, "observed_at": text, "not_top_3_pages": False}
    return {
        "rank": parse_int(match.group("rank")),
        "page": parse_int(match.group("page")),
        "position": parse_int(match.group("position")),
        "page_total": parse_int(match.group("total")),
        "observed_at": match.group("observed_at"),
        "not_top_3_pages": False,
    }


def parse_primary_traffic_keywords(asin: str, blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    text = ""
    for block in blocks:
        if block.get("name") == "main_relation":
            text = clean_text(block.get("text", ""))
            break
    if not text:
        return []
    text = re.sub(r"^主要流量词\s+收起\s+流量词\s+流量占比\s+流量词类型\s+自然排名\s+广告排名\s+", "", text)
    text = re.sub(r"\s+点击查看全部流量词.*$", "", text)
    row_pattern = re.compile(
        r"(?P<keyword>[A-Za-z0-9][A-Za-z0-9 '&+,\-/]+?)\s+"
        r"(?P<translation>[\u4e00-\u9fff][^0-9%]*?)\s+"
        r"(?P<share>[0-9]+(?:\.[0-9]+)?)%\s+"
        r"(?P<rest>.*?)(?=(?:[A-Za-z0-9][A-Za-z0-9 '&+,\-/]+?\s+[\u4e00-\u9fff][^0-9%]*?\s+[0-9]+(?:\.[0-9]+)?%)|$)"
    )
    rank_pattern = re.compile(r"(?:\d+\s+第\d+页,\d+/\d+\s+[^ ]+排名|前3页无排名)")
    rows: list[dict[str, Any]] = []
    for index, match in enumerate(row_pattern.finditer(text), start=1):
        rest = clean_text(match.group("rest"))
        rank_entries = rank_pattern.findall(rest)
        organic = parse_rank_entry(rank_entries[0] if rank_entries else "")
        ad = parse_rank_entry(rank_entries[1] if len(rank_entries) > 1 else "")
        rows.append(
            {
                "asin": asin,
                "position": index,
                "keyword": clean_text(match.group("keyword")),
                "keyword_translation": clean_text(match.group("translation")),
                "traffic_share": float(match.group("share")) / 100,
                "is_primary_traffic_keyword": "主要流量词" in rest,
                "is_natural_search_keyword": "自然搜索词" in rest,
                "is_brand_ad_keyword": "品牌广告词" in rest,
                "is_ad_keyword": "广告词" in rest or "广告流量词" in rest,
                "organic_rank": organic["rank"],
                "organic_page": organic["page"],
                "organic_position": organic["position"],
                "organic_page_total": organic["page_total"],
                "organic_observed_at": organic["observed_at"],
                "organic_not_top_3_pages": organic["not_top_3_pages"],
                "ad_rank": ad["rank"],
                "ad_page": ad["page"],
                "ad_position": ad["position"],
                "ad_page_total": ad["page_total"],
                "ad_observed_at": ad["observed_at"],
                "ad_not_top_3_pages": ad["not_top_3_pages"],
            }
        )
    return rows


def parse_inventory_offers(asin: str, blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    text = ""
    for block in blocks:
        if block.get("name") == "inventory":
            text = clean_text(block.get("text", ""))
            break
    if not text:
        return []
    text = re.sub(r"^卖家精灵-库存监控\s+剩余库存\s+", "", text)
    text = re.sub(r"\s+\d+\s+查看详细\s*$", "", text)
    offer_pattern = re.compile(r"(?P<quantity>[0-9,]+)\s+(?P<price>\$[0-9,]+(?:\.[0-9]+)?)\s+(?P<seller>.*?)(?=\s+[0-9,]+\s+\$|$)")
    rows = []
    for index, match in enumerate(offer_pattern.finditer(text), start=1):
        currency, price = parse_money(match.group("price"))
        rows.append(
            {
                "asin": asin,
                "offer_index": index,
                "inventory_quantity": parse_int(match.group("quantity")),
                "price": price,
                "currency": currency,
                "seller": clean_text(match.group("seller")),
            }
        )
    return rows


def parse_traffic_source_products(asin: str, rows: list[list[str]]) -> list[dict[str, Any]]:
    count_fields = [
        "all_traffic_keywords",
        "natural_search_keywords",
        "ac_recommendation_keywords",
        "er_recommendation_keywords",
        "four_star_recommendation_keywords",
        "hr_recommendation_keywords",
        "sp_ad_keywords",
        "video_ad_keywords",
        "brand_ad_keywords",
    ]
    product_pattern = re.compile(
        r"(?P<title>.*?)\s+变体数:(?P<variations>[0-9,]+)\s+"
        r"Color:\s*(?P<color>.*?)\s+\|\s+Size:\s*(?P<size>.*?)\s+"
        r"(?P<price>\$[0-9,]+(?:\.[0-9]+)?)\s+"
        r"(?P<rating>[0-9.]+)/(?P<rating_count>[0-9,]+)\s+"
        r"(?P<related_asin>[A-Z0-9]{10})$"
    )
    parsed = []
    for row in rows:
        product_idx = next((idx for idx, cell in enumerate(row) if re.search(r"\b[A-Z0-9]{10}\b", cell) and "$" in cell), -1)
        if product_idx < 0:
            continue
        product_info = clean_text(row[product_idx])
        match = product_pattern.search(product_info)
        if not match:
            continue
        source_modes = clean_text(row[product_idx + 1]) if product_idx + 1 < len(row) else ""
        count_values = row[product_idx + 2 : product_idx + 2 + len(count_fields)]
        currency, price = parse_money(match.group("price"))
        record: dict[str, Any] = {
            "input_asin": asin,
            "related_asin": match.group("related_asin"),
            "title": clean_text(match.group("title")),
            "variation_count": parse_int(match.group("variations")),
            "color": clean_text(match.group("color")),
            "size": clean_text(match.group("size")),
            "price": price,
            "currency": currency,
            "rating": parse_float(match.group("rating")),
            "rating_count": parse_int(match.group("rating_count")),
            "has_natural_search_source": "自然搜索" in source_modes,
            "has_amazon_recommendation_source": "亚马逊推荐" in source_modes,
            "has_ppc_ad_source": "PPC广告" in source_modes,
        }
        for field, value in zip(count_fields, count_values):
            record[field] = parse_int(value)
        parsed.append(record)
    return parsed


def parse_distribution(value: str) -> tuple[float | None, float | None]:
    natural = re.search(r"自然[:：]\s*([0-9.]+%)", value or "")
    ad = re.search(r"广告[:：]\s*([0-9.]+%)", value or "")
    return (
        parse_percent(natural.group(1)) if natural else None,
        parse_percent(ad.group(1)) if ad else None,
    )


def parse_keyword_reverse_rows(asin: str, rows: list[list[str]]) -> list[dict[str, Any]]:
    parsed = []
    for row in rows:
        if len(row) < 18 or not parse_int(row[0]):
            continue
        keyword, is_ac_from_keyword_cell, translation = parse_keyword_cell(row[1])
        traffic_values = split_numbers(row[2])
        monthly_search_values = split_numbers(row[9])
        purchase_values = split_numbers(row[12])
        impression_values = split_numbers(row[13])
        supply_values = split_numbers(row[14])
        aba_values = split_numbers(row[16])
        bid_values = split_numbers(row[17])
        natural_share, ad_share = parse_distribution(row[4])
        organic = parse_rank_entry(row[5])
        ad = parse_rank_entry(row[6])
        type_text = clean_text(row[3])
        tag_text = clean_text(row[2])
        parsed.append(
            {
                "asin": asin,
                "rank": parse_int(row[0]),
                "keyword": keyword,
                "keyword_translation": translation,
                "traffic_share": parse_percent(traffic_values[0]) if len(traffic_values) > 0 else None,
                "traffic_value": parse_int(traffic_values[1]) if len(traffic_values) > 1 else None,
                "is_primary_traffic_keyword": "主要流量词" in tag_text,
                "is_precise_keyword": "精准流量词" in tag_text,
                "is_precise_long_tail_keyword": "精准长尾词" in tag_text,
                "is_natural_search_keyword": "自然搜索词" in type_text,
                "is_amazon_choice_keyword": is_ac_from_keyword_cell or "AC推荐词" in type_text,
                "is_brand_ad_keyword": "品牌广告词" in type_text,
                "is_video_ad_keyword": "视频广告词" in type_text,
                "is_sp_ad_keyword": "SP广告词" in type_text,
                "natural_traffic_share": natural_share,
                "ad_traffic_share": ad_share,
                "organic_rank": organic["rank"],
                "organic_page": organic["page"],
                "organic_position": organic["position"],
                "organic_page_total": organic["page_total"],
                "organic_observed_at": organic["observed_at"],
                "organic_not_top_3_pages": organic["not_top_3_pages"],
                "ad_rank": ad["rank"],
                "ad_page": ad["page"],
                "ad_position": ad["position"],
                "ad_page_total": ad["page_total"],
                "ad_observed_at": ad["observed_at"],
                "ad_not_top_3_pages": ad["not_top_3_pages"],
                "aba_week_rank": parse_int(row[8]),
                "monthly_searches": parse_int(monthly_search_values[0]) if len(monthly_search_values) > 0 else None,
                "monthly_searches_aux": parse_int(monthly_search_values[1]) if len(monthly_search_values) > 1 else None,
                "spr": parse_int(row[10]),
                "title_density": parse_int(row[11]),
                "monthly_purchases": parse_int(purchase_values[0]) if len(purchase_values) > 0 else None,
                "purchase_rate": parse_percent(purchase_values[1]) if len(purchase_values) > 1 else None,
                "impressions": parse_int(impression_values[0]) if len(impression_values) > 0 else None,
                "clicks": parse_int(impression_values[1]) if len(impression_values) > 1 else None,
                "supply_demand_ratio": parse_float(supply_values[0]) if len(supply_values) > 0 else None,
                "products": parse_int(supply_values[1]) if len(supply_values) > 1 else None,
                "ad_products": parse_int(row[15]),
                "aba_click_share": parse_percent(aba_values[0]) if len(aba_values) > 0 else None,
                "aba_conversion_share": parse_percent(aba_values[1]) if len(aba_values) > 1 else None,
                "ppc_bid_low": parse_money(bid_values[0])[1] if len(bid_values) > 0 else None,
                "ppc_bid_exact": parse_money(bid_values[1])[1] if len(bid_values) > 1 else None,
                "ppc_bid_high": parse_money(bid_values[2])[1] if len(bid_values) > 2 else None,
            }
        )
    return parsed


def parse_keyword_cell(value: str) -> tuple[str, bool, str]:
    text = clean_text(value)
    is_ac = bool(re.search(r"\bAC\b", text))
    text = re.sub(r"\bAC\b", "", text).strip()
    match = re.search(r"^(?P<keyword>.*?)(?:\s+(?P<translation>[\u4e00-\u9fff].*))?$", text)
    if not match:
        return text, is_ac, ""
    return clean_text(match.group("keyword")), is_ac, clean_text(match.group("translation") or "")


def split_numbers(value: str) -> list[str]:
    return re.findall(r"\$?[0-9,]+(?:\.[0-9]+)?%?|\([^)]+\)", value or "")


def parse_keyword_miner_rows(seed_keyword: str, rows: list[list[str]]) -> list[dict[str, Any]]:
    parsed = []
    for row in rows:
        if len(row) < 16 or not parse_int(row[0]):
            continue
        keyword, is_ac, translation = parse_keyword_cell(row[1])
        relevancy_values = split_numbers(row[5])
        search_values = split_numbers(row[6])
        purchase_values = split_numbers(row[7])
        impression_values = split_numbers(row[8])
        supply_values = split_numbers(row[11])
        aba_share_values = split_numbers(row[13])
        bid_values = split_numbers(row[14])
        market_values = split_numbers(row[15])
        record = {
            "seed_keyword": seed_keyword,
            "rank": parse_int(row[0]),
            "keyword": keyword,
            "keyword_translation": translation,
            "is_amazon_choice_recommended": is_ac,
            "categories_text": clean_text(row[3]),
            "relevancy": parse_float(relevancy_values[0]) if len(relevancy_values) > 0 else None,
            "relevancy_aux": parse_float(relevancy_values[1]) if len(relevancy_values) > 1 else None,
            "monthly_searches": parse_int(search_values[0]) if len(search_values) > 0 else None,
            "monthly_searches_aux": parse_int(search_values[1]) if len(search_values) > 1 else None,
            "monthly_purchases": parse_int(purchase_values[0]) if len(purchase_values) > 0 else None,
            "purchase_rate": parse_percent(purchase_values[1]) if len(purchase_values) > 1 else None,
            "impressions": parse_int(impression_values[0]) if len(impression_values) > 0 else None,
            "clicks": parse_int(impression_values[1]) if len(impression_values) > 1 else None,
            "spr": parse_int(row[9]),
            "title_density": parse_int(row[10]),
            "supply_demand_ratio": parse_float(supply_values[0]) if len(supply_values) > 0 else None,
            "products": parse_int(supply_values[1]) if len(supply_values) > 1 else None,
            "ad_products": parse_int(row[12]),
            "aba_click_share": parse_percent(aba_share_values[0]) if len(aba_share_values) > 0 else None,
            "aba_conversion_share": parse_percent(aba_share_values[1]) if len(aba_share_values) > 1 else None,
            "ppc_bid_low": parse_money(bid_values[0])[1] if len(bid_values) > 0 else None,
            "ppc_bid_exact": parse_money(bid_values[1])[1] if len(bid_values) > 1 else None,
            "ppc_bid_high": parse_money(bid_values[2])[1] if len(bid_values) > 2 else None,
            "average_price": parse_money(market_values[0])[1] if len(market_values) > 0 else None,
            "average_rating_count": parse_int(market_values[1]) if len(market_values) > 1 else None,
            "average_rating": parse_float(market_values[2]) if len(market_values) > 2 else None,
        }
        parsed.append(record)
    return parsed


def wait_for_amazon_plugin(timeout: int) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    last: dict[str, Any] = {}
    js = r"""
(() => {
  const quick = document.querySelector("#seller-sprite-extension-quick-view-listing-page");
  const detail = document.querySelector("#seller-sprite-extension-quick-view-listing");
  const relation = document.querySelector("#seller-sprite-extension-main-relation");
  const text = [quick, detail, relation].filter(Boolean).map(el => el.innerText || el.textContent || "").join("\n");
  return JSON.stringify({
    url: location.href,
    title: document.title,
    hasQuick: Boolean(quick),
    hasDetail: Boolean(detail),
    hasRelation: Boolean(relation),
    textLength: text.length,
    hasSalesData: /近30天销量|销售额|BSR|FBA费用|主要流量词/.test(text)
  });
})()
"""
    while time.monotonic() < deadline:
        last = eval_js(AMAZON_SESSION, js, timeout=20)
        if isinstance(last, str):
            last = {"raw": last}
        if last.get("hasSalesData"):
            return last
        time.sleep(2)
    raise TimeoutError(f"Timed out waiting for SellerSprite Amazon extension data. Last state: {last}")


def extract_amazon_extension_blocks() -> dict[str, Any]:
    selectors_json = json.dumps(PLUGIN_SELECTORS, ensure_ascii=False)
    js = f"""
(() => {{
  const selectors = {selectors_json};
  const clean = s => (s || "").replace(/\\s+/g, " ").trim();
  const visible = el => {{
    const style = getComputedStyle(el);
    const rect = el.getBoundingClientRect();
    return style.display !== "none" && style.visibility !== "hidden" && Number(style.opacity) !== 0 && rect.width > 0 && rect.height > 0;
  }};
  const blocks = Object.entries(selectors).map(([name, selector]) => {{
    const el = document.querySelector(selector);
    if (!el) return {{name, selector, exists: false, visible: false, text: ""}};
    const rect = el.getBoundingClientRect();
    return {{
      name,
      selector,
      exists: true,
      visible: visible(el),
      rect: {{x: Math.round(rect.x), y: Math.round(rect.y), width: Math.round(rect.width), height: Math.round(rect.height)}},
      text: clean(el.innerText || el.textContent || "")
    }};
  }});
  return JSON.stringify({{
    url: location.href,
    title: document.title,
    extractedAt: new Date().toISOString(),
    blocks
  }});
}})()
"""
    data = eval_js(AMAZON_SESSION, js, timeout=45)
    if isinstance(data, str):
        data = json.loads(data)
    return data


def click_sellersprite_query_button() -> None:
    js = r"""
(() => {
  const visible = el => {
    const style = getComputedStyle(el);
    const rect = el.getBoundingClientRect();
    return style.display !== "none" && style.visibility !== "hidden" && Number(style.opacity) !== 0 && rect.width > 0 && rect.height > 0;
  };
  const button = [...document.querySelectorAll("button")]
    .find(el => visible(el) && /立即查询/.test(el.innerText || el.textContent || ""));
  if (!button) return JSON.stringify({clicked: false});
  button.click();
  return JSON.stringify({clicked: true, text: (button.innerText || "").trim()});
})()
"""
    result = eval_js(SELLERSPRITE_SESSION, js, timeout=20)
    if isinstance(result, str):
        result = json.loads(result)
    if not result.get("clicked"):
        raise RuntimeError(f"Could not click SellerSprite query button: {result}")


def wait_for_visible_table_rows(min_rows: int, timeout: int) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    last: dict[str, Any] = {}
    js = r"""
(() => {
  const clean = s => (s || "").replace(/\s+/g, " ").trim();
  const visible = el => {
    const style = getComputedStyle(el);
    const rect = el.getBoundingClientRect();
    return style.display !== "none" && style.visibility !== "hidden" && Number(style.opacity) !== 0 && rect.width > 0 && rect.height > 0;
  };
  const table = [...document.querySelectorAll(".el-table")].find(visible) || document.querySelector(".el-table");
  const rows = table ? [...table.querySelectorAll(".el-table__body-wrapper tbody tr")]
    .map(tr => [...tr.querySelectorAll("td")].map(td => clean(td.innerText)))
    .filter(r => r.some(Boolean)) : [];
  return JSON.stringify({url: location.href, title: document.title, rowCount: rows.length});
})()
"""
    while time.monotonic() < deadline:
        last = eval_js(SELLERSPRITE_SESSION, js, timeout=20)
        if isinstance(last, str):
            last = {"raw": last}
        if int(last.get("rowCount", 0)) >= min_rows:
            return last
        time.sleep(2)
    raise TimeoutError(f"Timed out waiting for visible table rows. Last state: {last}")


def extract_visible_sellersprite_table() -> dict[str, Any]:
    js = r"""
(() => {
  const clean = s => (s || "").replace(/\s+/g, " ").trim();
  const visible = el => {
    const style = getComputedStyle(el);
    const rect = el.getBoundingClientRect();
    return style.display !== "none" && style.visibility !== "hidden" && Number(style.opacity) !== 0 && rect.width > 0 && rect.height > 0;
  };
  const table = [...document.querySelectorAll(".el-table")].find(visible) || document.querySelector(".el-table");
  if (!table) return JSON.stringify({error: "No visible .el-table found", url: location.href, title: document.title});
  const headers = [...table.querySelectorAll(".el-table__header-wrapper th")]
    .map(th => clean(th.innerText))
    .filter(Boolean);
  const bodyRows = [...table.querySelectorAll(".el-table__body tbody tr")]
    .map(tr => [...tr.querySelectorAll("td")].map(td => clean(td.innerText)));
  const fixedRows = [...table.querySelectorAll(".el-table__fixed .el-table__fixed-body-wrapper tbody tr")]
    .map(tr => [...tr.querySelectorAll("td")].map(td => clean(td.innerText)));
  const rows = Array.from({length: Math.max(bodyRows.length, fixedRows.length)}, (_, idx) => {
    const body = bodyRows[idx] || [];
    const fixed = fixedRows[idx] || [];
    const maxLen = Math.max(body.length, fixed.length);
    return Array.from({length: maxLen}, (_, cellIdx) => fixed[cellIdx] || body[cellIdx] || "");
  }).filter(r => r.some(Boolean));
  return JSON.stringify({
    url: location.href,
    title: document.title,
    extractedAt: new Date().toISOString(),
    headers,
    rows,
    rowCount: rows.length
  });
})()
"""
    data = eval_js(SELLERSPRITE_SESSION, js, timeout=45)
    if isinstance(data, str):
        data = json.loads(data)
    if data.get("error"):
        raise RuntimeError(data["error"])
    return data


def query_sellersprite_traffic_source(asin: str, timeout: int) -> dict[str, Any]:
    run_opencli(SELLERSPRITE_SESSION, "open", SELLERSPRITE_REVERSING_URL, timeout=80)
    time.sleep(4)
    run_opencli(
        SELLERSPRITE_SESSION,
        "fill",
        'input[placeholder="输入父(子)体ASIN(最多20个)或关键词，多个ASIN以逗号区分"]',
        asin,
        timeout=30,
    )
    click_sellersprite_query_button()
    wait_for_visible_table_rows(1, timeout)
    return extract_visible_sellersprite_table()


def query_sellersprite_keyword_reverse(asin: str, timeout: int) -> dict[str, Any]:
    run_opencli(SELLERSPRITE_SESSION, "open", SELLERSPRITE_KEYWORD_REVERSE_URL, timeout=80)
    time.sleep(4)
    run_opencli(
        SELLERSPRITE_SESSION,
        "fill",
        'input[placeholder="请输入单个ASIN或产品链接，如: B00FLYWNYQ"]',
        asin,
        timeout=30,
    )
    click_sellersprite_query_button()
    wait_for_visible_table_rows(1, timeout)
    return extract_visible_sellersprite_table()


def query_keyword_miner(keyword: str, timeout: int, min_rows: int) -> dict[str, Any]:
    run_opencli(SELLERSPRITE_SESSION, "open", SELLERSPRITE_KEYWORD_MINER_URL, timeout=80)
    time.sleep(3)
    run_opencli(
        SELLERSPRITE_SESSION,
        "fill",
        'input[placeholder="输入关键词，如: flashlight"]',
        keyword,
        timeout=30,
    )
    click_sellersprite_query_button()
    wait_for_visible_table_rows(min_rows, timeout)
    return extract_visible_sellersprite_table()


def write_sheet(ws, rows: list[dict[str, Any]], columns: list[str]) -> None:
    ws.append(columns)
    for row in rows:
        ws.append([row.get(column, "") for column in columns])


def format_workbook(wb: Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            sample = [str(ws.cell(row=r, column=col_idx).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
            width = min(max(max((len(value) for value in sample), default=8) + 2, 10), 60)
            ws.column_dimensions[letter].width = width


def write_outputs(payload: dict[str, Any], out_dir: Path, asin: str) -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_path = out_dir / f"sellersprite_opencli_asin_keywords_{asin}_{stamp}.json"
    xlsx_path = out_dir / f"sellersprite_opencli_asin_keywords_{asin}_{stamp}.xlsx"
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    wb = Workbook()
    wb.remove(wb.active)
    sheets = [
        (
            "ASIN_Summary",
            payload["asin_summary"],
            ["asin", "field", "value", "unit", "source"],
        ),
        (
            "Primary_Traffic_Keywords",
            payload["primary_traffic_keywords"],
            [
                "asin",
                "position",
                "keyword",
                "keyword_translation",
                "traffic_share",
                "is_primary_traffic_keyword",
                "is_natural_search_keyword",
                "is_brand_ad_keyword",
                "is_ad_keyword",
                "organic_rank",
                "organic_page",
                "organic_position",
                "organic_page_total",
                "organic_observed_at",
                "organic_not_top_3_pages",
                "ad_rank",
                "ad_page",
                "ad_position",
                "ad_page_total",
                "ad_observed_at",
                "ad_not_top_3_pages",
            ],
        ),
        (
            "Inventory_Offers",
            payload["inventory_offers"],
            ["asin", "offer_index", "inventory_quantity", "price", "currency", "seller"],
        ),
        (
            "Traffic_Source_Products",
            payload["traffic_source_products"],
            [
                "input_asin",
                "related_asin",
                "title",
                "variation_count",
                "color",
                "size",
                "price",
                "currency",
                "rating",
                "rating_count",
                "has_natural_search_source",
                "has_amazon_recommendation_source",
                "has_ppc_ad_source",
                "all_traffic_keywords",
                "natural_search_keywords",
                "ac_recommendation_keywords",
                "er_recommendation_keywords",
                "four_star_recommendation_keywords",
                "hr_recommendation_keywords",
                "sp_ad_keywords",
                "video_ad_keywords",
                "brand_ad_keywords",
            ],
        ),
        (
            "Keyword_Reverse_Rows",
            payload["keyword_reverse_rows"],
            [
                "asin",
                "rank",
                "keyword",
                "keyword_translation",
                "traffic_share",
                "traffic_value",
                "is_primary_traffic_keyword",
                "is_precise_keyword",
                "is_precise_long_tail_keyword",
                "is_natural_search_keyword",
                "is_amazon_choice_keyword",
                "is_brand_ad_keyword",
                "is_video_ad_keyword",
                "is_sp_ad_keyword",
                "natural_traffic_share",
                "ad_traffic_share",
                "organic_rank",
                "organic_page",
                "organic_position",
                "organic_page_total",
                "organic_observed_at",
                "organic_not_top_3_pages",
                "ad_rank",
                "ad_page",
                "ad_position",
                "ad_page_total",
                "ad_observed_at",
                "ad_not_top_3_pages",
                "aba_week_rank",
                "monthly_searches",
                "monthly_searches_aux",
                "spr",
                "title_density",
                "monthly_purchases",
                "purchase_rate",
                "impressions",
                "clicks",
                "supply_demand_ratio",
                "products",
                "ad_products",
                "aba_click_share",
                "aba_conversion_share",
                "ppc_bid_low",
                "ppc_bid_exact",
                "ppc_bid_high",
            ],
        ),
        (
            "Keyword_Miner_Rows",
            payload["keyword_miner_rows"],
            [
                "seed_keyword",
                "rank",
                "keyword",
                "keyword_translation",
                "is_amazon_choice_recommended",
                "categories_text",
                "relevancy",
                "relevancy_aux",
                "monthly_searches",
                "monthly_searches_aux",
                "monthly_purchases",
                "purchase_rate",
                "impressions",
                "clicks",
                "spr",
                "title_density",
                "supply_demand_ratio",
                "products",
                "ad_products",
                "aba_click_share",
                "aba_conversion_share",
                "ppc_bid_low",
                "ppc_bid_exact",
                "ppc_bid_high",
                "average_price",
                "average_rating_count",
                "average_rating",
            ],
        ),
        (
            "Run_Metadata",
            payload["run_metadata"],
            ["field", "value"],
        ),
    ]
    for sheet_name, rows, columns in sheets:
        ws = wb.create_sheet(sheet_name)
        write_sheet(ws, rows, columns)
    format_workbook(wb)
    wb.save(xlsx_path)
    return json_path, xlsx_path


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Collect granular ASIN + keyword SellerSprite data through OpenCLI-controlled Chrome."
    )
    parser.add_argument("--asin", help="Target Amazon ASIN. If omitted, inferred from --amazon-url.")
    parser.add_argument("--amazon-url", help="Amazon product URL. Defaults to https://www.amazon.com/dp/{asin}.")
    parser.add_argument("--out-dir", default="data/sellersprite_opencli", help="Output directory.")
    parser.add_argument("--wait-timeout", type=int, default=75)
    parser.add_argument(
        "--keyword-miner-top",
        type=int,
        default=0,
        help="Optionally run Keyword Miner for the top N primary traffic keywords.",
    )
    parser.add_argument("--keyword-miner-min-rows", type=int, default=10)
    parser.add_argument("--skip-traffic-source", action="store_true", help="Skip SellerSprite /v3/reversing query.")
    parser.add_argument("--skip-keyword-reverse", action="store_true", help="Skip SellerSprite /v3/keyword-reverse query.")
    parser.add_argument("--skip-amazon-extension", action="store_true", help="Skip Amazon product page extension scrape.")
    args = parser.parse_args()

    if not args.asin and not args.amazon_url:
        parser.error("Provide --asin or --amazon-url")
    asin = args.asin or infer_asin(args.amazon_url or "")
    amazon_url = args.amazon_url or f"https://www.amazon.com/dp/{asin}"

    amazon_extension: dict[str, Any] = {"blocks": []}
    asin_summary: list[dict[str, Any]] = []
    primary_keywords: list[dict[str, Any]] = []
    inventory_offers: list[dict[str, Any]] = []
    traffic_source: dict[str, Any] = {"rows": []}
    traffic_source_products: list[dict[str, Any]] = []
    keyword_reverse: dict[str, Any] = {"rows": []}
    keyword_reverse_rows: list[dict[str, Any]] = []
    keyword_miner_tables: list[dict[str, Any]] = []
    keyword_miner_rows: list[dict[str, Any]] = []

    if not args.skip_amazon_extension:
        run_opencli(AMAZON_SESSION, "open", amazon_url, timeout=80)
        wait_for_amazon_plugin(args.wait_timeout)
        amazon_extension = extract_amazon_extension_blocks()
        asin_summary = parse_summary_metrics(asin, amazon_extension.get("blocks", []))
        primary_keywords = parse_primary_traffic_keywords(asin, amazon_extension.get("blocks", []))
        inventory_offers = parse_inventory_offers(asin, amazon_extension.get("blocks", []))

    if not args.skip_traffic_source:
        traffic_source = query_sellersprite_traffic_source(asin, args.wait_timeout)
        traffic_source_products = parse_traffic_source_products(asin, traffic_source.get("rows", []))

    if not args.skip_keyword_reverse:
        keyword_reverse = query_sellersprite_keyword_reverse(asin, args.wait_timeout)
        keyword_reverse_rows = parse_keyword_reverse_rows(asin, keyword_reverse.get("rows", []))

    keyword_seed_rows = primary_keywords if primary_keywords else keyword_reverse_rows
    for keyword_row in keyword_seed_rows[: max(args.keyword_miner_top, 0)]:
        keyword = keyword_row.get("keyword", "")
        if not keyword:
            continue
        table = query_keyword_miner(keyword, args.wait_timeout, args.keyword_miner_min_rows)
        table["seedKeyword"] = keyword
        keyword_miner_tables.append(table)
        keyword_miner_rows.extend(parse_keyword_miner_rows(keyword, table.get("rows", [])))

    payload = {
        "asin": asin,
        "amazon_url": amazon_url,
        "extracted_at": datetime.now().isoformat(timespec="seconds"),
        "asin_summary": asin_summary,
        "primary_traffic_keywords": primary_keywords,
        "inventory_offers": inventory_offers,
        "traffic_source_products": traffic_source_products,
        "keyword_reverse_rows": keyword_reverse_rows,
        "keyword_miner_rows": keyword_miner_rows,
        "run_metadata": [
            {"field": "source", "value": "OpenCLI-controlled Chrome with SellerSprite web and Chrome extension"},
            {"field": "asin", "value": asin},
            {"field": "amazon_url", "value": amazon_url},
            {"field": "amazon_extension_url", "value": amazon_extension.get("url", "")},
            {"field": "traffic_source_url", "value": traffic_source.get("url", "")},
            {"field": "keyword_reverse_url", "value": keyword_reverse.get("url", "")},
            {"field": "keyword_miner_tables", "value": len(keyword_miner_tables)},
            {
                "field": "amazon_extension_primary_keywords_available",
                "value": bool(primary_keywords),
            },
            {"field": "privacy_note", "value": "Excel stores granular business fields only; raw cookies/account headers/network bodies are not saved."},
            {"field": "boundary", "value": "This is visible web/plugin data, not SellerSprite MCP/API parity."},
        ],
        "raw_audit": {
            "amazon_extension": amazon_extension,
            "traffic_source": traffic_source,
            "keyword_reverse": keyword_reverse,
            "keyword_miner_tables": keyword_miner_tables,
        },
    }
    json_path, xlsx_path = write_outputs(payload, Path(args.out_dir), asin)
    print(
        json.dumps(
            {
                "asin": asin,
                "summaryRows": len(asin_summary),
                "primaryTrafficKeywordRows": len(primary_keywords),
                "inventoryOfferRows": len(inventory_offers),
                "trafficSourceProductRows": len(traffic_source_products),
                "keywordReverseRows": len(keyword_reverse_rows),
                "keywordMinerRows": len(keyword_miner_rows),
                "json": str(json_path.resolve()),
                "xlsx": str(xlsx_path.resolve()),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
